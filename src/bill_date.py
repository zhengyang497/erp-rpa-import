# -*- coding: utf-8 -*-
"""从 V2 导入模板读取统一账单日，并计算查询用的前一工作日。"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

# 持仓模板 / 成交模板列名
_DATE_HEADERS = ("持仓日期", "交易日", "交易日期", "账单日")


def _as_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s.replace("/", "-")[:10]).date()
    except ValueError:
        return None


def _is_workday(d: date) -> bool:
    """中国法定工作日（含调休）；库未覆盖该年时回退为周一至周五。"""
    try:
        from chinese_calendar import is_workday

        return bool(is_workday(d))
    except Exception:
        return d.weekday() < 5


def previous_workday(value: date | str) -> str:
    """账单日的前一个工作日，返回 YYYY-MM-DD。"""
    d = _as_date(value)
    if d is None:
        raise ValueError(f"无效日期: {value!r}")
    cur = d - timedelta(days=1)
    for _ in range(370):
        if _is_workday(cur):
            return cur.isoformat()
        cur -= timedelta(days=1)
    raise RuntimeError(f"无法计算 {d.isoformat()} 的前一工作日")


def read_bill_date(path: Path) -> str:
    """
    读取模板中的统一账单日，返回 YYYY-MM-DD。
    持仓列「持仓日期」、成交列「交易日」；要求非空行同一日。
    """
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise RuntimeError(f"空模板无表头: {path}")
        headers = [str(c).strip() if c is not None else "" for c in header]
        col = None
        for name in _DATE_HEADERS:
            if name in headers:
                col = headers.index(name)
                break
        if col is None:
            raise RuntimeError(
                f"未找到日期列 {_DATE_HEADERS}: {path.name} headers={headers[:12]}"
            )

        found: set[date] = set()
        for row in rows:
            if col >= len(row):
                continue
            d = _as_date(row[col])
            if d is not None:
                found.add(d)
        if not found:
            raise RuntimeError(f"模板无有效账单日: {path}")
        if len(found) > 1:
            raise RuntimeError(f"模板账单日不唯一: {sorted(found)} @ {path.name}")
        return next(iter(found)).isoformat()
    finally:
        wb.close()
